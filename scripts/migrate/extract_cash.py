"""
Cash book.xlsx :: 'Consolidated Petty Cash' -> CashTransaction rows.
Cross-checked against the duplicate copy of the same sheet embedded in
'Charlie Dairy 2026 - Cleaned.xlsx' (row count + final running balance).
The 'Heads' column (a loose secondary tag) has no dedicated field in the
schema for Phase 1, so it's folded into the remark text.
"""
import json
import os
import openpyxl
from common import SOURCE_DIR, OUTPUT_DIR, to_iso, to_float, clean_str, ReconLog

COL = dict(account=1, date=2, month=3, time=4, remark=5, heads=6, party=7,
           category=8, mode=9, entered_by=10, cash_in=11, cash_out=12,
           balance=13, project_land=14)


def parse_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        if row[COL["date"]] is None:
            continue
        rows.append(row)
    wb.close()
    return rows


def run(recon: ReconLog):
    primary_path = os.path.join(SOURCE_DIR, "Cash book.xlsx")
    rows = parse_sheet(primary_path, "Consolidated Petty Cash")

    transactions = []
    skipped = 0
    for row in rows:
        date_iso = to_iso(row[COL["date"]])
        if not date_iso:
            skipped += 1
            continue
        remark = clean_str(row[COL["remark"]])
        heads = clean_str(row[COL["heads"]])
        combined_remark = f"[{heads}] {remark}" if heads and remark else (remark or heads)
        mode_raw = (clean_str(row[COL["mode"]]) or "cash").lower()
        transactions.append({
            "date": date_iso,
            "time": clean_str(row[COL["time"]]),
            "account": clean_str(row[COL["account"]]),
            "party": clean_str(row[COL["party"]]),
            "category": clean_str(row[COL["category"]]) or "Uncategorized",
            "mode": "BANK" if "bank" in mode_raw else "CASH",
            "amountIn": to_float(row[COL["cash_in"]]),
            "amountOut": to_float(row[COL["cash_out"]]),
            "enteredBy": clean_str(row[COL["entered_by"]]),
            "projectLand": clean_str(row[COL["project_land"]]),
            "remark": combined_remark,
        })

    with open(os.path.join(OUTPUT_DIR, "cash_transactions.json"), "w", encoding="utf-8") as f:
        json.dump(transactions, f, indent=2)

    dates = [t["date"] for t in transactions]
    recon.note(f"[cash] Parsed {len(transactions)} cash transactions from Cash book.xlsx ({min(dates)} to {max(dates)})" if dates else "[cash] No transactions parsed")
    if skipped:
        recon.note(f"[cash] Skipped {skipped} rows with unparsable dates")

    # Cross-check against the duplicate copy in the Cleaned workbook.
    dup_path = os.path.join(SOURCE_DIR, "Charlie Dairy 2026 - Cleaned.xlsx")
    try:
        dup_rows = parse_sheet(dup_path, "Consolidated Petty Cash")
        primary_last_balance = to_float(rows[-1][COL["balance"]]) if rows else None
        dup_last_balance = to_float(dup_rows[-1][COL["balance"]]) if dup_rows else None
        recon.note(f"[cash] CROSS-CHECK: Cash book.xlsx has {len(rows)} rows, last balance {primary_last_balance}. "
                   f"Cleaned.xlsx copy has {len(dup_rows)} rows, last balance {dup_last_balance}.")
        if primary_last_balance != dup_last_balance or len(rows) != len(dup_rows):
            recon.note("[cash] ⚠ MISMATCH between the two copies — using Cash book.xlsx as source of truth. "
                       "Verify with the user which file is actually the most recently updated before trusting cash figures.")
        else:
            recon.note("[cash] Copies match — no discrepancy found.")
    except Exception as e:
        recon.note(f"[cash] Could not cross-check against Cleaned.xlsx copy: {e}")

    return transactions


if __name__ == "__main__":
    log = ReconLog()
    run(log)
    log.write(os.path.join(OUTPUT_DIR, "reconciliation_report.md"))
