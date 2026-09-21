"""
Cow master.xlsx :: 'Daily Log' -> one row per milking (cow, shift, litres).
Rows with Shift == 'Calve' are calving-event markers, not milkings; they are
skipped here and counted for the reconciliation report.
"""
import json
import os
import openpyxl
from common import SOURCE_DIR, OUTPUT_DIR, to_iso, to_float, clean_str, ReconLog

SHIFT_MAP = {"morning": "MORNING", "afternoon": "AFTERNOON", "evening": "EVENING"}


def run(recon: ReconLog):
    path = os.path.join(SOURCE_DIR, "Cow master.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Daily Log"]

    records = []
    skipped_calve = 0
    skipped_bad = 0
    unknown_shift = set()

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        date_raw, tag_raw, shift_raw, litres_raw = row[0], row[1], row[2], row[3]
        if date_raw is None or tag_raw is None:
            continue
        shift_key = (clean_str(shift_raw) or "").lower()
        if shift_key == "calve":
            skipped_calve += 1
            continue
        shift = SHIFT_MAP.get(shift_key)
        if not shift:
            if shift_raw:
                unknown_shift.add(shift_raw)
            skipped_bad += 1
            continue
        date_iso = to_iso(date_raw)
        if not date_iso:
            skipped_bad += 1
            continue
        records.append({
            "cowTag": clean_str(tag_raw),
            "date": date_iso,
            "shift": shift,
            "litres": to_float(litres_raw),
        })

    wb.close()

    with open(os.path.join(OUTPUT_DIR, "milking_records.json"), "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2)

    dates = [r["date"] for r in records]
    recon.note(f"[milking] Parsed {len(records)} milking records ({min(dates)} to {max(dates)})" if dates else "[milking] No records parsed")
    recon.note(f"[milking] Skipped {skipped_calve} calving-event marker rows, {skipped_bad} malformed rows")
    if unknown_shift:
        recon.note(f"[milking] Unrecognized shift values (skipped, review manually): {unknown_shift}")
    return records


if __name__ == "__main__":
    log = ReconLog()
    run(log)
    log.write(os.path.join(OUTPUT_DIR, "reconciliation_report.md"))
