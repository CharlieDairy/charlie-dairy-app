"""
Capital.xlsx :: 'Capital ' has TWO tables that both look like partner ledgers:
  - Columns D-K ('PHASE 3 - Cash Flow - BANK'): only 12 rows, Mar 2024 only.
  - Columns N-U ('Date/From/To...Investment'): 150 rows spanning 2021-2026,
    tagged per entry with a venture (Fattening, Dairy, LOAN 22, LOAN 23-1/2/3,
    Phase 3.1/3.2, LOAN 24-1/2, LOAN 25-1).
Cross-checking the two showed overlapping partners/amounts around Mar 2024
under the 'Phase 3.1' venture tag (e.g. Abid Rs 1,140,000 appears in both) —
so only the N-U table is imported, to avoid double-counting. The D-K table is
skipped entirely and flagged for the user to confirm.

Only entries tagged venture == 'Dairy' are specific to Charlie Dairy Farm;
other ventures (Fattening, various loan tranches) are shared across the same
partner group but a different business line. All are imported (partners
overlap and loans may fund dairy operations too) but the venture tag is kept
on every row so the app can filter to Dairy-only if the user wants that.

BL.xlsx :: 'BL Accounts' has four side-by-side GL-style sub-ledgers
(Agri Land Lease, Capital Investment - Maaz Khan, Agri Cost Recovery,
Bank Account - Abid). These are account ledgers, not partner names, but are
modeled the same way (running Dr/Cr ledger) for Phase 1 — 'partner' holds the
account/ledger name instead. Flagged for the user to confirm this
simplification is acceptable, or split into a proper accounts model later.

Both sources are cross-checked against their duplicate copies embedded in
'Charlie Dairy 2026 - Cleaned.xlsx'.
"""
import json
import os
import openpyxl
from common import SOURCE_DIR, OUTPUT_DIR, to_iso, to_float, clean_str, ReconLog

BL_BLOCKS = {1: "Agri Land Lease", 7: "Capital Investment - Maaz Khan",
             13: "Agri Cost Recovery Acc from Charlie", 19: "Bank Account - Abid"}


def parse_capital_sheet(path, sheet_name="Capital "):
    """Parses the N-U 'Date/From/To...Investment' table (cols idx 13-20)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        if len(row) <= 20 or row[13] is None:
            continue
        date_iso = to_iso(row[13])
        if not date_iso:
            continue
        credit, debit = to_float(row[16]), to_float(row[17])
        rows.append({
            "date": date_iso,
            "partner": clean_str(row[14]) or "Unknown",
            "description": clean_str(row[15]) or "",
            "bankAccount": None,
            "debit": debit,
            "credit": credit,
            "type": "CONTRIBUTION" if credit > 0 else ("WITHDRAWAL" if debit > 0 else "OTHER"),
            "venture": clean_str(row[20]),
        })
    wb.close()
    return rows


def parse_bl_sheet(path, sheet_name="BL Accounts"):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        for start, label in BL_BLOCKS.items():
            date_cell = row[start]
            if date_cell is None:
                continue
            date_iso = to_iso(date_cell)
            if not date_iso:
                continue
            debit, credit = to_float(row[start + 2]), to_float(row[start + 3])
            rows.append({
                "date": date_iso, "partner": label,
                "description": clean_str(row[start + 1]) or "",
                "bankAccount": None,
                "debit": debit, "credit": credit,
                "type": "CONTRIBUTION" if credit > 0 else ("WITHDRAWAL" if debit > 0 else "OTHER"),
                "venture": None,
            })
    wb.close()
    return rows


def run(recon: ReconLog):
    cap_rows = parse_capital_sheet(os.path.join(SOURCE_DIR, "Capital.xlsx"))
    bl_rows = parse_bl_sheet(os.path.join(SOURCE_DIR, "BL.xlsx"))
    entries = cap_rows + bl_rows

    with open(os.path.join(OUTPUT_DIR, "capital_entries.json"), "w", encoding="utf-8") as f:
        json.dump(entries, f, indent=2)

    ventures = {}
    for r in cap_rows:
        ventures[r["venture"]] = ventures.get(r["venture"], 0) + 1
    dairy_count = ventures.get("Dairy", 0)

    recon.note(f"[capital] Capital.xlsx partner ledger (N-U table): {len(cap_rows)} entries across ventures {ventures}")
    recon.note(f"[capital] Only {dairy_count} entries are tagged 'Dairy' specifically — the rest are other ventures "
               f"(Fattening) or loan tranches shared by the same partner group. All are imported with the venture "
               f"tag preserved; confirm with the user whether Charlie Dairy Farm reporting should filter to 'Dairy' only.")
    recon.note("[capital] ⚠ Skipped the smaller 'PHASE 3 - Cash Flow - BANK' table (cols D-K, 12 rows) — its Mar 2024 "
               "entries appear to duplicate amounts already in the main ledger under venture 'Phase 3.1' "
               "(e.g. Abid Rs 1,140,000 appears in both). Verify with the user this table isn't recording something distinct.")
    recon.note(f"[capital] BL.xlsx sub-ledgers ({', '.join(BL_BLOCKS.values())}): {len(bl_rows)} entries")
    recon.note("[capital] ⚠ Note: BL.xlsx entries are GL-style account ledgers, not partner names — stored in the "
               "same 'partner' field as a simplification for Phase 1. Confirm with the user whether these should be "
               "split into a distinct accounts model later.")

    dup_path = os.path.join(SOURCE_DIR, "Charlie Dairy 2026 - Cleaned.xlsx")
    try:
        dup_cap = parse_capital_sheet(dup_path, "Capital")
        recon.note(f"[capital] CROSS-CHECK Capital: Capital.xlsx has {len(cap_rows)} rows, Cleaned.xlsx copy has {len(dup_cap)} rows.")
        if len(cap_rows) != len(dup_cap):
            recon.note("[capital] ⚠ MISMATCH in Capital ledger row counts — using Capital.xlsx as source of truth. Verify with the user.")
    except Exception as e:
        recon.note(f"[capital] Could not cross-check Capital sheet against Cleaned.xlsx: {e}")

    try:
        dup_bl = parse_bl_sheet(dup_path, "BL Accounts")
        recon.note(f"[capital] CROSS-CHECK BL Accounts: BL.xlsx has {len(bl_rows)} rows, Cleaned.xlsx copy has {len(dup_bl)} rows.")
        if len(bl_rows) != len(dup_bl):
            recon.note("[capital] ⚠ MISMATCH in BL Accounts row counts — using BL.xlsx as source of truth. Verify with the user.")
    except Exception as e:
        recon.note(f"[capital] Could not cross-check BL Accounts against Cleaned.xlsx: {e}")

    return entries


if __name__ == "__main__":
    log = ReconLog()
    run(log)
    log.write(os.path.join(OUTPUT_DIR, "reconciliation_report.md"))
