"""
Charlie Dairy 2026 - Cleaned.xlsx :: 'Assets List' -> Asset rows.
The sheet has two valuation snapshots (a 'Current Value - Feb 25' column and
a later one dated 21/4/2025); the more recent (21/4/2025) is used as the
current value, falling back to the Feb '25 figure if the later one is blank.
"""
import datetime
import json
import os
import openpyxl
from common import SOURCE_DIR, OUTPUT_DIR, to_float, clean_str, ReconLog


def run(recon: ReconLog):
    path = os.path.join(SOURCE_DIR, "Charlie Dairy 2026 - Cleaned.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Assets List"]

    assets = []
    fallback_used = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        asset_class, details = clean_str(row[1]), clean_str(row[2])
        if not asset_class and not details:
            continue
        recent_value = row[9] if len(row) > 9 else None
        valuation_date = "2025-04-21"
        if recent_value in (None, ""):
            recent_value = row[7]
            valuation_date = "2025-02-28"
            fallback_used += 1
        assets.append({
            "assetClass": asset_class or "Uncategorized",
            "details": details or "",
            "qty": to_float(row[3], default=0),
            "value": to_float(row[4], default=0),
            "depreciationPct": to_float(row[5], default=0),
            "yearLived": int(to_float(row[6], default=0)),
            "currentValue": to_float(recent_value, default=0),
            "valuationDate": valuation_date,
        })

    wb.close()

    with open(os.path.join(OUTPUT_DIR, "assets.json"), "w", encoding="utf-8") as f:
        json.dump(assets, f, indent=2)

    total_value = sum(a["currentValue"] for a in assets)
    recon.note(f"[assets] Parsed {len(assets)} asset lines, total current value Rs {total_value:,.0f}")
    if fallback_used:
        recon.note(f"[assets] {fallback_used} lines had no 21/4/2025 value and fell back to the Feb '25 snapshot")
    return assets


if __name__ == "__main__":
    log = ReconLog()
    run(log)
    log.write(os.path.join(OUTPUT_DIR, "reconciliation_report.md"))
