"""
Cow master.xlsx :: 'Feed Cost' -> FeedTransaction rows.
Sheet layout is 4 side-by-side feed-type blocks (Silage, Wenda, Turi, Fodder),
each with Inward/Outward/Balance/Rate/Cost columns per date. The first data
row is an 'Opening' balance (no date) and is captured as a note, not a
transaction, since it has no in/out movement to record.
"""
import json
import os
import openpyxl
from common import SOURCE_DIR, OUTPUT_DIR, to_iso, to_float, clean_str, ReconLog

BLOCKS = {"Silage": 2, "Wenda": 8, "Turi": 14, "Fodder": 20}


def run(recon: ReconLog):
    path = os.path.join(SOURCE_DIR, "Cow master.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Feed Cost"]

    transactions = []
    openings = {}
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        date_cell = row[1]
        if date_cell is None:
            continue
        if clean_str(date_cell) == "Opening":
            for feed_type, start in BLOCKS.items():
                openings[feed_type] = to_float(row[start])
            continue
        date_iso = to_iso(date_cell)
        if not date_iso:
            continue
        for feed_type, start in BLOCKS.items():
            inward, outward, balance, rate, cost = (
                row[start], row[start + 1], row[start + 2], row[start + 3], row[start + 4]
            )
            inward_v = to_float(inward, default=None) if inward not in (None, "") else None
            outward_v = to_float(outward, default=None) if outward not in (None, "") else None
            if inward_v:
                transactions.append({
                    "date": date_iso, "feedType": feed_type, "direction": "IN",
                    "quantity": inward_v, "rate": None, "cost": None,
                })
            if outward_v:
                transactions.append({
                    "date": date_iso, "feedType": feed_type, "direction": "OUT",
                    "quantity": outward_v, "rate": to_float(rate, default=None),
                    "cost": to_float(cost, default=None),
                })

    wb.close()

    with open(os.path.join(OUTPUT_DIR, "feed_transactions.json"), "w", encoding="utf-8") as f:
        json.dump(transactions, f, indent=2)

    dates = [t["date"] for t in transactions]
    recon.note(f"[feed] Parsed {len(transactions)} feed in/out transactions across {len(BLOCKS)} feed types "
               f"({min(dates)} to {max(dates)})" if dates else "[feed] No transactions parsed")
    recon.note(f"[feed] Opening balances captured (not migrated as transactions, informational only): {openings}")
    return transactions


if __name__ == "__main__":
    log = ReconLog()
    run(log)
    log.write(os.path.join(OUTPUT_DIR, "reconciliation_report.md"))
